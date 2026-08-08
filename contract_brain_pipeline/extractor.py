#!/usr/bin/env python3
"""Deterministic source extraction used by Contract Brain.

Design goals copied from the proven CONOPS evidence workflow:

* hash source bytes before analysis;
* use native text extraction first;
* invoke OCR only when a PDF has no native text;
* preserve exact source locators;
* keep source rows normalized instead of repeating full text on relationships;
* make archive handling bounded and explicit.

This module does not perform semantic promotion.  It only produces source evidence
rows that later Contract Brain modules may review and relate.
"""

from __future__ import annotations

import hashlib
import json
import re
import shutil
import subprocess
import tempfile
import zipfile
from dataclasses import dataclass, field
from io import BytesIO
from pathlib import Path, PurePosixPath
from typing import Any

import fitz  # PyMuPDF
from docx import Document
from openpyxl import load_workbook
from pptx import Presentation

PARSER_VERSION = "contract-brain-extractor/0.1.0"
MAX_ARCHIVE_MEMBERS = 200
MAX_ARCHIVE_MEMBER_BYTES = 100 * 1024 * 1024
MAX_TEXT_BYTES = 32 * 1024 * 1024
PLAIN_TEXT_EXTS = {
    ".txt", ".csv", ".tsv", ".md", ".json", ".xml", ".html", ".htm",
    ".yaml", ".yml", ".rtf", ".log", ".ini", ".cfg", ".conf",
}
SUPPORTED_EXTS = {".pdf", ".docx", ".pptx", ".xlsx", ".zip"} | PLAIN_TEXT_EXTS


@dataclass(frozen=True)
class LogicalRow:
    locator: str
    text: str
    row_type: str
    page_number: int | None = None
    geometry: dict[str, float] | None = None
    sort_index: int = 0
    member_path: str = ""


@dataclass
class ExtractionResult:
    rows: list[LogicalRow] = field(default_factory=list)
    status: str = "complete"
    mode: str = ""
    page_count: int | None = None
    ocr_page_count: int = 0
    error: str = ""
    metadata: dict[str, Any] = field(default_factory=dict)


def clean_text(value: object) -> str:
    return " ".join(str(value or "").replace("\x00", "").split())


def sha256_bytes(payload: bytes) -> str:
    return hashlib.sha256(payload).hexdigest()


def source_row_id(source_sha256: str, row: LogicalRow, parser_version: str = PARSER_VERSION) -> str:
    material = json.dumps(
        [parser_version, source_sha256, row.member_path, row.locator, row.row_type, row.text],
        ensure_ascii=False,
        separators=(",", ":"),
    ).encode("utf-8")
    return hashlib.sha256(material).hexdigest()


def _pdf_native_rows(payload: bytes) -> tuple[list[LogicalRow], int]:
    document = fitz.open(stream=payload, filetype="pdf")
    rows: list[LogicalRow] = []
    order = 0
    try:
        page_count = document.page_count
        for page_index in range(page_count):
            page = document.load_page(page_index)
            for block_index, block in enumerate(page.get_text("blocks", sort=True), start=1):
                if len(block) < 5:
                    continue
                x0, y0, x1, y1, raw_text = block[:5]
                block_type = int(block[6]) if len(block) > 6 and isinstance(block[6], (int, float)) else 0
                if block_type != 0:
                    continue
                lines = [clean_text(line) for line in str(raw_text or "").splitlines()]
                lines = [line for line in lines if line]
                for line_no, line in enumerate(lines, start=1):
                    order += 1
                    rows.append(
                        LogicalRow(
                            locator=f"page {page_index + 1} block {block_index} line {line_no}",
                            text=line,
                            row_type="pdf_native",
                            page_number=page_index + 1,
                            geometry={"x0": float(x0), "y0": float(y0), "x1": float(x1), "y1": float(y1)},
                            sort_index=order,
                        )
                    )
        return rows, page_count
    finally:
        document.close()


def _pdf_ocr_rows(payload: bytes) -> tuple[list[LogicalRow], int]:
    tesseract = shutil.which("tesseract")
    if not tesseract:
        raise RuntimeError("Tesseract is not installed; zero-native-text PDF requires OCR fallback")

    document = fitz.open(stream=payload, filetype="pdf")
    rows: list[LogicalRow] = []
    order = 0
    try:
        with tempfile.TemporaryDirectory(prefix="contract-brain-ocr-") as tmp:
            tmp_path = Path(tmp)
            matrix = fitz.Matrix(200 / 72, 200 / 72)
            for page_index in range(document.page_count):
                page = document.load_page(page_index)
                pixmap = page.get_pixmap(matrix=matrix, colorspace=fitz.csGRAY, alpha=False)
                image_path = tmp_path / f"page-{page_index + 1:04d}.png"
                pixmap.save(str(image_path))
                completed = subprocess.run(
                    [tesseract, str(image_path), "stdout", "-l", "eng", "--psm", "3"],
                    stdout=subprocess.PIPE,
                    stderr=subprocess.PIPE,
                    timeout=120,
                    check=False,
                )
                if completed.returncode != 0:
                    stderr = completed.stderr.decode("utf-8", errors="replace")
                    raise RuntimeError(
                        f"Tesseract failed on page {page_index + 1} rc={completed.returncode}: {clean_text(stderr)[:800]}"
                    )
                line_number = 0
                for raw_line in completed.stdout.decode("utf-8", errors="replace").splitlines():
                    line = clean_text(raw_line)
                    if not line:
                        continue
                    line_number += 1
                    order += 1
                    rows.append(
                        LogicalRow(
                            locator=f"page {page_index + 1} OCR line {line_number}",
                            text=line,
                            row_type="pdf_ocr",
                            page_number=page_index + 1,
                            sort_index=order,
                        )
                    )
        return rows, document.page_count
    finally:
        document.close()


def _docx_rows(payload: bytes) -> list[LogicalRow]:
    document = Document(BytesIO(payload))
    rows: list[LogicalRow] = []
    order = 0
    for index, paragraph in enumerate(document.paragraphs, start=1):
        value = clean_text(paragraph.text)
        if value:
            order += 1
            rows.append(LogicalRow(f"paragraph {index}", value, "docx_paragraph", sort_index=order))
    for table_index, table in enumerate(document.tables, start=1):
        for row_index, table_row in enumerate(table.rows, start=1):
            cells = [clean_text(cell.text) for cell in table_row.cells]
            value = " | ".join(f"C{i}={cell}" for i, cell in enumerate(cells, start=1) if cell)
            if value:
                order += 1
                rows.append(LogicalRow(f"table {table_index} row {row_index}", value, "docx_table", sort_index=order))
    return rows


def _pptx_rows(payload: bytes) -> list[LogicalRow]:
    presentation = Presentation(BytesIO(payload))
    rows: list[LogicalRow] = []
    order = 0
    for slide_no, slide in enumerate(presentation.slides, start=1):
        for shape_index, shape in enumerate(slide.shapes, start=1):
            if getattr(shape, "has_text_frame", False):
                value = clean_text(shape.text)
                if value:
                    order += 1
                    rows.append(LogicalRow(f"slide {slide_no} shape {shape_index}", value, "pptx_text", page_number=slide_no, sort_index=order))
            if getattr(shape, "has_table", False):
                for row_index, table_row in enumerate(shape.table.rows, start=1):
                    cells = [clean_text(cell.text) for cell in table_row.cells]
                    value = " | ".join(f"C{i}={cell}" for i, cell in enumerate(cells, start=1) if cell)
                    if value:
                        order += 1
                        rows.append(LogicalRow(f"slide {slide_no} shape {shape_index} table row {row_index}", value, "pptx_table", page_number=slide_no, sort_index=order))
    return rows


def _xlsx_rows(payload: bytes) -> list[LogicalRow]:
    workbook = load_workbook(BytesIO(payload), read_only=True, data_only=True)
    rows: list[LogicalRow] = []
    order = 0
    try:
        for sheet in workbook.worksheets:
            for row_index, values in enumerate(sheet.iter_rows(values_only=True), start=1):
                parts = []
                for column_index, value in enumerate(values, start=1):
                    cleaned = clean_text(value)
                    if cleaned:
                        parts.append(f"C{column_index}={cleaned}")
                if parts:
                    order += 1
                    rows.append(LogicalRow(f"{sheet.title}!row {row_index}", " | ".join(parts), "xlsx_row", sort_index=order))
    finally:
        workbook.close()
    return rows


def _plain_text_rows(payload: bytes, suffix: str) -> list[LogicalRow]:
    if len(payload) > MAX_TEXT_BYTES:
        raise RuntimeError(f"text source exceeds bounded {MAX_TEXT_BYTES} byte limit")
    text = payload.decode("utf-8", errors="replace")
    rows: list[LogicalRow] = []
    order = 0
    for line_no, raw_line in enumerate(text.splitlines(), start=1):
        value = clean_text(raw_line)
        if value:
            order += 1
            rows.append(LogicalRow(f"line {line_no}", value, f"text_{suffix.lstrip('.') or 'plain'}", sort_index=order))
    return rows


def _prefix_member(rows: list[LogicalRow], member: str, offset: int) -> list[LogicalRow]:
    out: list[LogicalRow] = []
    for index, row in enumerate(rows, start=1):
        out.append(
            LogicalRow(
                locator=f"member {member} :: {row.locator}",
                text=row.text,
                row_type=row.row_type,
                page_number=row.page_number,
                geometry=row.geometry,
                sort_index=offset + index,
                member_path=member,
            )
        )
    return out


def _zip_rows(payload: bytes) -> ExtractionResult:
    rows: list[LogicalRow] = []
    archive_members: list[dict[str, Any]] = []
    with zipfile.ZipFile(BytesIO(payload)) as archive:
        members = [info for info in archive.infolist() if not info.is_dir()]
        if len(members) > MAX_ARCHIVE_MEMBERS:
            return ExtractionResult(status="review_required", mode="ZIP inventory only", error=f"archive has {len(members)} members; bounded limit is {MAX_ARCHIVE_MEMBERS}")
        for info in members:
            member_name = PurePosixPath(info.filename).as_posix()
            suffix = Path(member_name).suffix.lower()
            archive_members.append({"name": member_name, "size": info.file_size, "supported": suffix in SUPPORTED_EXTS and suffix != ".zip"})
            if info.file_size > MAX_ARCHIVE_MEMBER_BYTES or suffix not in SUPPORTED_EXTS or suffix == ".zip":
                continue
            member_payload = archive.read(info)
            result = extract_payload(member_name, member_payload, suffix=suffix, allow_archive=False)
            rows.extend(_prefix_member(result.rows, member_name, len(rows)))
    return ExtractionResult(
        rows=rows,
        status="complete" if rows else "review_required",
        mode="bounded one-level ZIP member extraction",
        metadata={"archive_members": archive_members},
    )


def extract_payload(name: str, payload: bytes, *, suffix: str | None = None, allow_archive: bool = True) -> ExtractionResult:
    suffix = (suffix or Path(name).suffix).lower()
    try:
        if suffix == ".pdf":
            native_rows, page_count = _pdf_native_rows(payload)
            if native_rows:
                return ExtractionResult(
                    rows=native_rows,
                    mode="PyMuPDF native text blocks; sort=True; OCR not invoked",
                    page_count=page_count,
                )
            ocr_rows, page_count = _pdf_ocr_rows(payload)
            if ocr_rows:
                return ExtractionResult(
                    rows=ocr_rows,
                    mode="PyMuPDF zero-native-text gate -> 200 DPI grayscale -> Tesseract eng psm 3",
                    page_count=page_count,
                    ocr_page_count=page_count,
                )
            return ExtractionResult(status="review_required", mode="native-first PDF extraction", page_count=page_count, error="native and OCR extraction returned zero text")
        if suffix == ".docx":
            return ExtractionResult(rows=_docx_rows(payload), mode="python-docx paragraphs and table rows")
        if suffix == ".pptx":
            rows = _pptx_rows(payload)
            pages = max((row.page_number or 0 for row in rows), default=0) or None
            return ExtractionResult(rows=rows, mode="python-pptx slide text and tables", page_count=pages)
        if suffix == ".xlsx":
            return ExtractionResult(rows=_xlsx_rows(payload), mode="openpyxl sparse row extraction")
        if suffix in PLAIN_TEXT_EXTS:
            return ExtractionResult(rows=_plain_text_rows(payload, suffix), mode="UTF-8 text line extraction")
        if suffix == ".zip" and allow_archive:
            return _zip_rows(payload)
        return ExtractionResult(status="review_required", mode="unsupported source type", error=f"unsupported extension: {suffix or '(none)'}")
    except Exception as exc:
        return ExtractionResult(status="failed", mode="deterministic source extraction", error=str(exc))
